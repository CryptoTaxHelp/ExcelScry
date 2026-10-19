"""Generate a sample workbook for trying out ExcelScry.

The point of this file is to exercise the cases that break naive viewers:
several sheets, a table that does not start at A1, a time-of-day column
(Excel stores those as a fraction of a day, which renders as a bogus 1899 date
if you are not careful), an integer too large for a float, a blank row inside
the used range, markup-looking text, and sheets that are wider and taller than
the viewport. The money, date and percentage columns carry their own number
formats, so the grid has to show `$23,988.00`, `15/01/2026` and `35.0%` rather
than the raw values.

Run from the repository root:  python scripts/make_sample.py
"""

from __future__ import annotations

import datetime as dt
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = pathlib.Path(__file__).resolve().parent.parent / "samples" / "ExcelScry-sample.xlsx"


def add_report(book: Workbook) -> None:
    """A normal-looking report whose table starts at C6."""
    ws = book.active
    ws.title = "Sales report"

    first_row, first_col = 6, 3  # C6
    headers = ["Region", "Product", "Units", "Unit price", "Revenue", "Ordered", "Notes"]
    for offset, header in enumerate(headers):
        cell = ws.cell(row=first_row, column=first_col + offset, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")

    rows = [
        ("North", "Widget", 1200, 19.99, 23988.0, dt.date(2026, 1, 15), "repeat customer"),
        ("South", "Widget", 940, 19.99, 18790.6, dt.date(2026, 2, 3), None),
        ("East", "Gadget", 0, 249.5, 0.0, dt.date(2026, 2, 20), "backordered"),
        ("West", "Gadget", 37, 249.5, 9231.5, dt.date(2026, 3, 1), "expedited"),
        (
            "North",
            "Doohickey",
            18446744073709551,  # beyond 2^53: a float round trip would corrupt it
            0.5,
            9223372036854775.0,
            dt.date(2026, 3, 14),
            "large id",
        ),
        (None, None, None, None, None, None, None),  # blank row inside the range
        ("South", "Doohickey", 8, 0.1, 0.8, dt.date(2026, 4, 2), "rounding"),
        ("East", "Widget", 12, 19.99, 239.88, dt.date(2026, 4, 19), None),
        ("West", "Gadget", 3, 249.5, 748.5, dt.date(2026, 5, 7), "#N/A upstream"),
        ("North", "Widget", 5, 19.99, 99.95, dt.date(2026, 5, 30), "TRUE"),
    ]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            ws.cell(row=first_row + 1 + r, column=first_col + c, value=value)

    # Times of day and a margin, in their own columns.
    shift_col = first_col + len(headers) + 1
    ws.cell(row=first_row, column=shift_col, value="Start time").font = Font(bold=True)
    ws.cell(row=first_row, column=shift_col + 1, value="Margin").font = Font(bold=True)
    times = [dt.time(9, 30), dt.time(13, 5, 15), dt.time(23, 59, 59)]
    for r, value in enumerate(times):
        ws.cell(row=first_row + 1 + r, column=shift_col, value=value)
    for r, margin in enumerate((0.35, 0.28, 0.0, 0.42, 0.6, 0.15, 0.2, 0.3)):
        ws.cell(row=first_row + 1 + r, column=shift_col + 1, value=margin)

    # The formats Excel/Sheets would write. Without them the viewer would show
    # 23988, 46037 and 0.35 where the sheet shows $23,988.00, 15/01/2026 and
    # 35.0%.
    for r in range(first_row + 1, first_row + 1 + len(rows)):
        ws.cell(row=r, column=first_col + 2).number_format = "#,##0"
        ws.cell(row=r, column=first_col + 3).number_format = '"$"#,##0.00'
        ws.cell(row=r, column=first_col + 4).number_format = '"$"#,##0.00'
        ws.cell(row=r, column=first_col + 5).number_format = "dd/mm/yyyy"
    for r in range(first_row + 1, first_row + 1 + len(times)):
        ws.cell(row=r, column=shift_col).number_format = "hh:mm"
    for r in range(first_row + 1, first_row + 9):
        ws.cell(row=r, column=shift_col + 1).number_format = "0.0%"

    for column, width in zip("CDEFGHIJK", (10, 12, 8, 11, 12, 12, 18, 11, 9)):
        ws.column_dimensions[column].width = width


def add_wide(book: Workbook) -> None:
    """Wider than the viewport, so horizontal virtual scrolling is exercised."""
    ws = book.create_sheet("Wide sheet")
    for c in range(1, 61):
        ws.cell(row=1, column=c, value=f"col {c}")
    for r in range(2, 40):
        for c in range(1, 61):
            ws.cell(row=r, column=c, value=(r * c) % 97)


def add_tall(book: Workbook) -> None:
    """Taller than the viewport, so row virtualisation is exercised."""
    ws = book.create_sheet("Tall sheet")
    ws.cell(row=1, column=1, value="index")
    ws.cell(row=1, column=2, value="value")
    for r in range(2, 2002):
        ws.cell(row=r, column=1, value=r - 1)
        ws.cell(row=r, column=2, value=(r * 37) % 1000 / 4)


def add_escaping(book: Workbook) -> None:
    """Text that must be shown literally, never interpreted as markup."""
    ws = book.create_sheet("Unicode & escaping")
    values = [
        "Markup-looking text must not be interpreted:",
        "<script>alert('xss')</script>",
        'a "quoted" & <angled> value',
        "emoji ✅ and accents éöü",
        "=1+1 stays text when exported",
        "000123",
        True,
        False,
    ]
    for r, value in enumerate(values, start=1):
        ws.cell(row=r, column=1, value=value)
    ws.cell(row=1, column=2, value="multi\nline\ncell")
    ws.cell(row=2, column=2, value="a very long value that should be clipped in the grid and offered in full via the cell readout tooltip")


def main() -> None:
    book = Workbook()
    add_report(book)
    add_wide(book)
    add_tall(book)
    book.create_sheet("Empty sheet")
    add_escaping(book)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
